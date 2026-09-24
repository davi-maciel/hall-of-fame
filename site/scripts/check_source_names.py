#!/usr/bin/env python3
"""
Name-presence check for every corroboration URL, per (olympiad, year, person).

A person page lists every corroborating URL of an edition, but many of those
pages never mention a given person (medal-only lists, team-count news, another
student's profile). This script fetches each URL referenced by
<ol>/data/corroboration.json once, extracts its text (HTML stripped of tags,
PDFs via pypdf, diacritics-insensitive lowercase fold) and decides, for every
roster member of that (olympiad, year) — site/data/people.json `name` plus
`nameVariants` — whether the page names the person:

  full       the whole name appears (particles de/da/do/dos/das/e ignored,
             initials optional; a glued run such as "JoaoSilva" in PDF text
             also counts when it is long enough to be unambiguous)
  firstlast  first name and last non-suffix surname (Filho/Neto/Júnior ignored)
             appear together: adjacent, or with middle names / initials between
             them (known to the pool or not), or surname-first ("Timbó, Rafael")
  variant    the same, but one whole token is spelled one edit away
             ("Bernardo Peruzzo Trevisan" for ...Trevizan): every other token
             matches exactly, the differing one is >= 4 letters (>= 5 for a
             first name), and no token of the span belongs to a teammate
  surname    the last two significant tokens appear adjacent ("Studart Neto")
  short      the page prints a shortened form: own tokens in the person's order
             with at least one dropped ("Daniel Lima" for Daniel Lima Braga),
             the halves in two table cells ("Alexandrino | Davi"), or a
             fixed-width truncation ("Ulisses Fonsec"). Weakest level: it only
             counts when no teammate of that edition could be meant instead.
  attested   the source's text cannot be fetched or parsed at all (a social post
             behind a login wall, a diploma image), and the curation attests that
             it names the person: the source entry in <ol>/data/corroboration.json
             carries "names": ["<person-id>", ...] (person ids = the slugs in
             site/data/people.json). Weakest level of all — a curator's reading of
             the source, not this script's; it never overrides a real text match.

Sources that are not plain HTML are opened too: spreadsheets are read cell by
cell (one line per row, cells joined with " | "), PDFs without a text layer and
image sources are rendered and read with the macOS Vision OCR helper
(`ocr_vision.swift`, compiled on demand), a Yandex Disk public link is resolved
to its real download through the public API, and a URL whose live fetch fails
falls back to the newest Wayback snapshot (never replacing a live answer).

Matching is substring-based on a folded copy of the text (lowercase, no
diacritics, separators kept as "|"), so table cells glued together by PDF
extraction ("27Rafael TimbóBrazil3,5") still match; a match may not start in
the middle of a word, and only long enough names may run into the next cell.
Tokens of the person's teammates in that edition guard against reading a
list of names as one person.

Best level wins; anything else = not named. Pages that yield no usable text
(dead links, images, spreadsheets, JS shells, social login walls) get status
"error" / "status-only" with an empty found-set: unverifiable, not negative —
unless the source entry attests names, which is exactly what "names" is for.

Writes site/data/source_names.json
  {"generatedAt": ..., "byOlympiad": {ol: {year: {url:
      {"status": "text"|"status-only"|"error", "found": {personId: level}}}}}}
and tmp/source_names_report.md (per-dataset tallies plus the list of
participations that have no name-bearing source at all — the gap list).
Its name-expansion candidates skip the person ids recorded in site/scripts/aliases.json
`expansionsRejected` — forms already looked at and turned down (source typo, spelling-only).

Fetched text is cached under tmp/source_cache/<sha1(url)>.txt (+ .json meta),
so reruns only re-analyse. Requests are serialised per host with a small delay.

Run:  python3 site/scripts/check_source_names.py
      [--only imo,icho] [--retry-errors] [--refetch status-only,error]
      [--skip-hosts noic.com.br] [--workers 8] [--delay 1.0]
"""
import argparse
import bisect
import hashlib
import html
import io
import json
import re
import ssl
import subprocess
import sys
import threading
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from functools import lru_cache
from datetime import datetime, timezone
from pathlib import Path

SITE = Path(__file__).resolve().parent.parent
ROOT = SITE.parent
PEOPLE = SITE / "data" / "people.json"
OUT = SITE / "data" / "source_names.json"
CACHE = ROOT / "tmp" / "source_cache"
REPORT = ROOT / "tmp" / "source_names_report.md"
NOTES = ROOT / "tmp" / "gap_notes.json"  # optional, {"<ol> <year> <pid>": [class, note]}
OCR_SRC = SITE / "scripts" / "ocr_vision.swift"
OCR_DIR = ROOT / "tmp" / "ocr"
OCR_BIN = OCR_DIR / "ocr_vision"

sys.path.insert(0, str(SITE / "scripts"))
from build_people import DATASETS, slugify  # noqa: E402  (single source of truth for the dataset list and the slug rule)

UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
TIMEOUT = 25
MAX_BYTES = 8_000_000
MAX_DOC_BYTES = 96_000_000  # a scanned results PDF runs to tens of MB, and a truncated one
#                             cannot be opened at all — pages are still capped for HTML
MIN_TEXT = 200  # folded chars; a 200 page with less than this is a shell, not a document
PARTICLES = {"de", "da", "do", "dos", "das", "e"}
# a conjunction that the fold drops (or punctuates away) between two tokens separates
# two names: "Vilian Borchardt e Ronaldo Rodrigues" is a list, not one long name
CONJUNCTIONS = {"e", "and", "y"}
SUFFIXES = {"filho", "neto", "junior", "sobrinho"}
RANK = {"full": 5, "firstlast": 4, "variant": 3, "surname": 2, "short": 1, "attested": 0}
LEVELS = ["full", "firstlast", "variant", "surname", "short", "attested"]
# anonymous fetches of these hosts return a login wall / JS shell, never the post
NO_TEXT_HOSTS = {"instagram.com", "facebook.com", "twitter.com", "x.com", "linkedin.com", "youtube.com", "youtu.be", "tiktok.com"}
BINARY_CT = ("image/", "video/", "audio/", "zip", "msword", "vnd.", "octet-stream", "font/")
IMAGE_CT = ("image/jpeg", "image/png", "image/gif", "image/webp", "image/tiff", "image/bmp", "image/heic")
IMAGE_EXT = (".jpg", ".jpeg", ".png", ".gif", ".webp", ".tif", ".tiff", ".bmp", ".heic")
SHEET_CT = ("spreadsheetml", "ms-excel", "msexcel", "excel")
SHEET_EXT = (".xlsx", ".xlsm")
# public file hosts that serve a JS shell to a plain GET; their real file needs an API call
YANDEX_HOSTS = {"disk.yandex.ru", "disk.yandex.com", "disk.360.yandex.ru", "disk.360.yandex.com", "yadi.sk"}
YANDEX_API = "https://cloud-api.yandex.net/v1/disk/public/resources/download?public_key="
WAYBACK_CDX = "https://web.archive.org/cdx/search/cdx?url={}&output=json&fl=timestamp&filter=statuscode:200&limit=-3"
WAYBACK_AVAIL = "https://archive.org/wayback/available?url={}"
WAYBACK_RAW = "https://web.archive.org/web/{}id_/{}"  # id_ = the bytes as archived, no toolbar
ARCHIVE_TIMEOUT = 75  # the CDX index is slow for rarely-asked URLs
OCR_LANGS = "pt-BR,en-US,es-ES,ru-RU"
OCR_DPI = 200
OCR_MAX_PAGES = 40


# ----------------------------------------------------------------------------- text folding

BOUNDARY = set(",;:()[]{}|/\\•·–—\n\r\t")  # separators between list items, table cells, lines


def fold(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").lower()
    s = s.replace("'", "").replace("’", "").replace("`", "")
    return re.sub(r"[^a-z0-9]+", " ", s)


def tokens(s: str):
    return ["junior" if t == "jr" else t for t in fold(s).split()]


def fold_chars(raw: str):
    """Per-character fold that keeps raw offsets: letters/digits lowercased without
    diacritics, apostrophes dropped, separators -> '|', anything else -> ' '."""
    chars, cmap = [], []
    for idx, ch in enumerate(raw):
        if ch in "'’`":
            continue
        if ch in BOUNDARY:
            chars.append("|")
            cmap.append(idx)
        elif ch < "\x80":
            c = ch.lower()
            chars.append(c if ("a" <= c <= "z" or "0" <= c <= "9") else " ")
            cmap.append(idx)
        else:
            for c in unicodedata.normalize("NFKD", ch):
                if unicodedata.category(c) == "Mn":
                    continue
                c = c.lower()
                chars.append(c if ("a" <= c <= "z" or "0" <= c <= "9") else " ")
                cmap.append(idx)
    return chars, cmap


class Doc:
    """Folded document as one string: words separated by single spaces, '|' where
    the source had a separator, particles removed. Glued table cells
    ("27rafael timbobrazil3") stay glued — matching is substring-based."""

    def __init__(self, raw: str):
        self.raw = raw
        chars, cmap = fold_chars(raw)
        toks, spans = [], []  # token text, (raw start, raw end)
        i, n = 0, len(chars)
        while i < n:
            c = chars[i]
            if c == " ":
                i += 1
            elif c == "|":
                if toks and toks[-1] != "|":
                    toks.append("|")
                    spans.append((cmap[i], cmap[i] + 1))
                i += 1
            else:
                j = i
                while j < n and chars[j] != " " and chars[j] != "|":
                    j += 1
                t = "".join(chars[i:j])
                if t == "jr":
                    t = "junior"
                if t not in PARTICLES:
                    toks.append(t)
                    spans.append((cmap[i], cmap[j - 1] + 1))
                i = j
        if toks and toks[-1] == "|":
            toks.pop()
            spans.pop()
        self.toks, self.spans = toks, spans
        self.text = " ".join(toks)
        self.tpos, off = [], 0  # char offset of every token in self.text
        for t in toks:
            self.tpos.append(off)
            off += len(t) + 1
        self.squashed = "".join(t for t in toks if t != "|")
        self._pos = self._tokset = None

    @property
    def positions(self):
        """token -> [indices in self.toks] (separators excluded), built on demand."""
        if self._pos is None:
            p = defaultdict(list)
            for i, t in enumerate(self.toks):
                if t != "|":
                    p[t].append(i)
            self._pos = p
        return self._pos

    @property
    def tokset(self):
        if self._tokset is None:
            self._tokset = frozenset(self.positions)
        return self._tokset

    def token_at(self, i: int) -> str:
        return self.toks[i] if 0 <= i < len(self.toks) else ""

    def tok_index(self, pos: int) -> int:
        return bisect.bisect_right(self.tpos, pos) - 1

    def conj_between(self, k: int, m: int) -> bool:
        """True when the raw text between two neighbouring tokens is exactly a
        conjunction. "e" folds away with the particles, so "Borchardt e Ronaldo"
        reads as one run here although the page lists two people."""
        if not (0 <= k < len(self.spans) and 0 <= m < len(self.spans)):
            return False
        a, b = self.spans[min(k, m)][1], self.spans[max(k, m)][0]
        return a <= b and self.raw[a:b].strip().lower() in CONJUNCTIONS

    def neighbour(self, k: int, step: int) -> str:
        """The token next to index k (step -1 before, +1 after), or "" when a dropped
        conjunction stands between the two: it separates names, so whatever follows
        it is not hugging the span."""
        m = k + step
        if not (0 <= m < len(self.toks)) or self.conj_between(k, m):
            return ""
        return self.toks[m]

    def token_before(self, start: int) -> str:
        return self.neighbour(self.tok_index(start), -1)

    def token_after(self, end: int) -> str:
        return self.neighbour(self.tok_index(end - 1), 1)

    def raw_span(self, start: int, end: int) -> str:
        a = self.spans[self.tok_index(start)][0]
        b = self.spans[self.tok_index(end - 1)][1]
        return re.sub(r"\s+", " ", self.raw[a:b]).strip()


class NameSpec:
    def __init__(self, form: str):
        toks = tokens(form)
        self.sig = [t for t in toks if t not in PARTICLES] or toks      # "joao pedro t silva junior"
        self.core = [t for t in self.sig if len(t) > 1] or self.sig     # initials dropped
        nonsuf = [t for t in self.core if t not in SUFFIXES] or self.core
        self.first, self.last = nonsuf[0], nonsuf[-1]
        self.own = set(self.sig)
        self.middle = set(self.core) - {self.first, self.last}  # middle names + suffixes
        self.key = tuple(self.sig)


def _pat(words, glue_ok=False):
    """The words in sequence; nothing alphabetic may precede ("27rafael" is fine,
    "banana lima" is not); running into the next cell ("timbobrazil") only when allowed."""
    return r"(?<![a-z])" + re.escape(" ".join(words)) + ("" if glue_ok else r"(?![a-z])")


def _long(words):
    return sum(len(w) for w in words) >= 10


# --------------------------------------------------------------------- near-miss (variant) match
# Sources misspell names ("Trevisan" for Trevizan, "Almerida" for Almeida). One
# whole token one edit away is still a naming; two are a different person, so the
# rest of the name must match exactly, token for token, and nothing in the span
# may belong to a teammate. Matching is token-based here (no glued cells).

ALPHA = "abcdefghijklmnopqrstuvwxyz"
VAR_MIN = 4        # a surname / middle name may differ from this length on
VAR_MIN_FIRST = 5  # a first name only from this length on (Luiz/Luis stays a non-match)
VAR_GAP = 32       # middle-name budget between first and last, as in the firstlast branch


@lru_cache(maxsize=8192)
def near_forms(w: str) -> frozenset:
    """Every string exactly one edit away from w (substitution, insertion, deletion,
    transposition of adjacent letters — Damerau-Levenshtein 1); w itself excluded."""
    out = set()
    for i in range(len(w)):
        out.add(w[:i] + w[i + 1:])
        if i + 1 < len(w):
            out.add(w[:i] + w[i + 1] + w[i] + w[i + 2:])
        for c in ALPHA:
            out.add(w[:i] + c + w[i + 1:])
            out.add(w[:i] + c + w[i:])
    out.update(w + c for c in ALPHA)
    out.discard(w)
    return frozenset(out)


def _near_positions(doc: Doc, word: str, minlen: int):
    """Positions of doc tokens one edit away from `word` (never the word itself)."""
    if len(word) < minlen:
        return []
    return [i for w in near_forms(word) & doc.tokset for i in doc.positions[w]]


def _seq_variant(doc: Doc, seq, first, others):
    """seq matched token for token against consecutive doc tokens, exactly one off."""
    n, k_order = len(doc.toks), sorted(range(len(seq)), key=lambda k: -len(seq[k]))[:2]
    tried = set()
    for k in k_order:  # one of the two longest tokens is the exact one -> a valid anchor
        for p in doc.positions.get(seq[k], ()):
            s = p - k
            if s < 0 or s + len(seq) > n or s in tried:
                continue
            tried.add(s)
            run = doc.toks[s:s + len(seq)]
            if any(t == "|" or t in others for t in run):
                continue
            off = [q for q in range(len(seq)) if run[q] != seq[q]]
            if len(off) != 1:
                continue
            q = off[0]
            if len(seq[q]) < (VAR_MIN_FIRST if seq[q] == first else VAR_MIN):
                continue
            if run[q] not in near_forms(seq[q]):
                continue
            if doc.neighbour(s, -1) in others or doc.neighbour(s + len(seq) - 1, 1) in others:
                continue
            return s, s + len(seq) - 1
    return None


def _reach(doc: Doc, i: int, targets, allowed, others):
    """First target position after i, reached across middle names the pool already
    knows (plus initials). Unlike the firstlast branch this does not step over
    middle names unknown to the pool: with a misspelt surname, "Lucas H Morais"
    for Lucas Levy ... Moraes is evidence, "Lucas Henrique Morais" is a coin flip."""
    toks, budget, j = doc.toks, VAR_GAP, i + 1
    while j < len(toks):
        if j in targets and toks[j] not in others:
            return j
        t = toks[j]
        if t in others or not (t in allowed or (len(t) == 1 and t.isalpha())):
            return None
        budget -= len(t) + 1
        if budget < 0:
            return None
        j += 1
    return None


def _gap_variant(doc: Doc, sp: NameSpec, others):
    """first <own middle names> last with exactly one of first/last one edit off."""
    exact_first = doc.positions.get(sp.first, ())
    exact_last = set(doc.positions.get(sp.last, ()))
    allowed = sp.middle
    for starts, targets in (
        (exact_first, set(_near_positions(doc, sp.last, VAR_MIN))),
        (_near_positions(doc, sp.first, VAR_MIN_FIRST), exact_last),
    ):
        for i in starts:
            if doc.toks[i] in others or doc.neighbour(i, -1) in others:
                continue
            j = _reach(doc, i, targets, allowed, others)
            if j is not None and doc.neighbour(j, 1) not in others:
                return i, j
    return None


def variant_match(doc: Doc, sp: NameSpec, others):
    """-> (first token index, last token index) of a one-edit near miss, or None."""
    for seq in (sp.sig, sp.core):
        if len(seq) >= 2:
            hit = _seq_variant(doc, seq, sp.first, others)
            if hit:
                return hit
    if sp.first != sp.last:
        return _gap_variant(doc, sp, others)
    return None


# ----------------------------------------------------------------- short-form match
# Result tables and news often print a shortened name: the first given name plus
# some of the following tokens with the identifying surname dropped ("Daniel Lima"
# for Daniel Lima Braga), the two halves in separate table cells ("Alexandrino |
# Davi"), or a fixed-width truncation ("Ulisses Fonsec"). A shortened form is weak
# evidence on its own, so it counts only when it cannot be read as a teammate of
# that same edition (edition-scoped uniqueness) and does not hug a teammate's name.

SHORT_MIN_TOK = 2   # tokens in the span
SHORT_MIN_LEN = 8   # letters in the span
SHORT_SIDE = 4      # each side of a column split needs a token this long
TRUNC_MIN = 5       # letters of a truncated final token


def _idx(own, tok, k):
    try:
        return own.index(tok, k)
    except ValueError:
        return None


def _own_run(doc: Doc, i: int, own, k0: int):
    """Longest contiguous doc run from i whose tokens are own[k0:] in the name's
    order -> (last index, run tokens, next own index), or None."""
    m = _idx(own, doc.toks[i], k0)
    if m is None:
        return None
    run, k, j = [doc.toks[i]], m + 1, i
    while j + 1 < len(doc.toks):
        m = _idx(own, doc.toks[j + 1], k)
        if m is None:
            break
        run.append(doc.toks[j + 1])
        k, j = m + 1, j + 1
    return j, run, k


def _order_ok(own, run):
    k = -1
    for t in run:
        k = _idx(own, t, k + 1)
        if k is None:
            return False
    return True


def _letters(run):
    return sum(len(t) for t in run)


def _mates_ok(run, mates, trunc=None):
    """Edition-scoped uniqueness: no teammate's own tokens contain the whole span."""
    for mt in mates:
        if all(t in mt for t in run) and (trunc is None or any(t.startswith(trunc) for t in mt)):
            return False
    return True


def _edges_ok(doc: Doc, i: int, j: int, others):
    return doc.neighbour(i, -1) not in others and doc.neighbour(j, 1) not in others


def _alone(doc: Doc, i: int, j: int, run, pool):
    """No neighbouring token turns the span into somebody else's name: "João Victor"
    is not a short form of João Victor Soares Aleixo when the page reads
    "João Victor Teixeira Degelo" and that person exists in the same dataset."""
    for nb in (doc.neighbour(i, -1), doc.neighbour(j, 1)):
        if len(nb) < 3 or not nb.isalpha():
            continue
        if any(nb in q and all(t in q for t in run) for q in pool):
            return False
    return True


def short_match(doc: Doc, sp: NameSpec, others, mates, pool=()):
    """-> (first token index, last token index) of a shortened form, or None.
    Every token of the span is one of the person's own, so the "no teammate token"
    guard reduces to the uniqueness test plus the neighbour tests."""
    own, n = sp.sig, len(doc.toks)
    own_set = set(own)
    # 1. first given name + own tokens in order, at least one own token dropped
    k_first = own.index(sp.first)
    for i in doc.positions.get(sp.first, ()):
        got = _own_run(doc, i, own, k_first)
        if not got:
            continue
        j, run, _ = got
        if (len(run) >= SHORT_MIN_TOK and len(run) < len(own) and _letters(run) >= SHORT_MIN_LEN
                and _edges_ok(doc, i, j, others) and _mates_ok(run, mates) and _alone(doc, i, j, run, pool)):
            return i, j
    # 2. the two halves in separate cells, either order ("Alexandrino | Davi")
    for t in own_set:
        for q in doc.positions.get(t, ()):
            if doc.token_at(q + 1) != "|" or doc.token_at(q + 2) not in own_set:
                continue
            left, p = [], q
            while p >= 0 and doc.toks[p] in own_set and _order_ok(own, doc.toks[p:q + 1]):
                left, p = doc.toks[p:q + 1], p - 1
            right, e = [], q + 2
            while e < n and doc.toks[e] in own_set and _order_ok(own, doc.toks[q + 2:e + 1]):
                right, e = doc.toks[q + 2:e + 1], e + 1
            run = left + right
            if len(set(run)) != len(run) or len(run) < SHORT_MIN_TOK or _letters(run) < SHORT_MIN_LEN:
                continue
            if max(map(len, left)) < SHORT_SIDE or max(map(len, right)) < SHORT_SIDE:
                continue
            i, j = q - len(left) + 1, q + 1 + len(right)
            if _edges_ok(doc, i, j, others) and _mates_ok(run, mates) and _alone(doc, i, j, run, pool):
                return i, j
    # 3. fixed-width truncation of the last token ("Ulisses Fonsec |"), after at
    # least one exactly matched own token
    for t in own_set:
        if len(t) < SHORT_SIDE:
            continue
        for i in doc.positions.get(t, ()):
            got = _own_run(doc, i, own, 0)
            if not got:
                continue
            j, run, k = got
            cut = doc.token_at(j + 1)
            if len(cut) < TRUNC_MIN or cut in own_set or cut in others:
                continue
            if not any(own[m].startswith(cut) for m in range(k, len(own))):
                continue
            after = doc.token_at(j + 2)
            if after not in ("", "|") and not after[0].isdigit():
                continue
            if (_letters(run + [cut]) >= SHORT_MIN_LEN and _edges_ok(doc, i, j + 1, others)
                    and _mates_ok(run, mates, trunc=cut) and _alone(doc, i, j + 1, run, pool)):
                return i, j + 1
    return None


def match_level(doc: Doc, sp: NameSpec, others=frozenset(), mates=(), pool=()):
    """-> (level, longer_form, variant_span). `others` = name tokens of the person's
    teammates in that edition; `longer_form` is the page's text when it shows first +
    middle names unknown to the pool + last (a name-expansion candidate);
    `variant_span` is the raw page text of a `variant` match, for human review."""
    T = doc.text
    for seq in (sp.sig, sp.core):
        if re.search(_pat(seq, _long(seq)), T):
            return "full", None, None
    joined = "".join(sp.core)
    if len(sp.core) >= 2 and len(joined) >= 12 and joined in doc.squashed:
        return "full", None, None
    if sp.first != sp.last:
        pair = [sp.first, sp.last]
        if re.search(_pat(pair, _long(pair)), T):
            return "firstlast", None, None
        # first <middle names / initials, known to the pool or not> last: no separator,
        # no digit, <= 32 chars, and no token that belongs to a teammate
        gap = re.escape(sp.first) + r"(?![a-z]) ([a-z](?:[a-z ]{0,30}[a-z])?) " + re.escape(sp.last) + r"(?![a-z])"
        for m in re.finditer(r"(?<![a-z])" + gap, T):
            if not any(g in others for g in m.group(1).split()):
                return "firstlast", doc.raw_span(m.start(), m.end()), None
        # "first | last": the two halves in separate table cells ("Larissa | Lima" in a
        # nombre/apellido table), the mirror of the surname-first rule below and guarded
        # the same way — the span may not hug a teammate's name
        fwd = r"(?<![a-z])" + re.escape(sp.first) + r"(?![a-z]) \| " + re.escape(sp.last) + r"(?![a-z])"
        for m in re.finditer(fwd, T):
            if doc.token_before(m.start()) not in others and doc.token_after(m.end()) not in others:
                return "firstlast", None, None
        # surname-first: "last [own middle tokens] [,] first" — not when hugging teammate tokens
        own = sorted(sp.middle | {t for t in sp.sig if len(t) == 1})
        mid = (r"(?: (?:%s))*" % "|".join(map(re.escape, own))) if own else ""
        rev = r"(?<![a-z])" + re.escape(sp.last) + r"(?![a-z])" + mid + r"(?: \|)? " + re.escape(sp.first) + r"(?![a-z])"
        for m in re.finditer(rev, T):
            if doc.token_before(m.start()) not in others and doc.token_after(m.end()) not in others:
                return "firstlast", None, None
    hit = variant_match(doc, sp, others)
    if hit:
        i, j = hit
        return "variant", None, doc.raw_span(doc.tpos[i], doc.tpos[j] + len(doc.toks[j]))
    if len(sp.core) >= 2 and re.search(_pat(sp.core[-2:], _long(sp.core[-2:])), T):
        return "surname", None, None
    hit = short_match(doc, sp, others, mates, pool)
    if hit:
        i, j = hit
        return "short", None, doc.raw_span(doc.tpos[i], doc.tpos[j] + len(doc.toks[j]))
    return None, None, None


# ----------------------------------------------------------------------------- fetching

def fetch(url, _depth=0, timeout=None):
    req = urllib.request.Request(
        url,
        headers={  # full desktop-Chrome headers: some hosts (noic.com.br) 403/406 anything terser
            "User-Agent": UA,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,application/pdf,*/*;q=0.8",
            "Accept-Language": "pt-BR,pt;q=0.9,en;q=0.8",
            "Upgrade-Insecure-Requests": "1",
        },
    )
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE  # some .gov.br chains are broken; the content is what matters
    try:
        with urllib.request.urlopen(req, timeout=timeout or TIMEOUT, context=ctx) as r:
            ct = r.headers.get("Content-Type", "")
            cap = MAX_DOC_BYTES if any(x in ct.lower() for x in ("pdf",) + SHEET_CT) else MAX_BYTES
            return r.status, ct, r.read(cap), r.geturl()
    except urllib.error.HTTPError as e:
        # urllib < 3.11 does not auto-follow 308
        if e.code in (301, 302, 307, 308) and e.headers.get("Location") and _depth < 3:
            return fetch(urllib.parse.urljoin(url, e.headers["Location"]), _depth + 1, timeout)
        raise


class HostGate:
    """At most one in-flight request per host, spaced by `delay` seconds."""

    def __init__(self, delay):
        self.delay = delay
        self.mu = threading.Lock()
        self.locks = defaultdict(threading.Lock)
        self.last = {}

    def fetch(self, url, timeout=None):
        host = urllib.parse.urlparse(url).netloc.lower()
        with self.mu:
            lock = self.locks[host]
        with lock:
            wait = self.last.get(host, 0.0) + self.delay - time.monotonic()
            if wait > 0:
                time.sleep(wait)
            try:
                return fetch(url, timeout=timeout)
            finally:
                self.last[host] = time.monotonic()


# ------------------------------------------------------------------ OCR (macOS Vision)
# Result PDFs are often scans and some sources are only a photo of the podium list.
# `ocr_vision.swift` turns a bitmap into text with the system's text recogniser and
# keeps the table layout (one line per row, cells joined with " | "), so the column
# rules below see the same shape they see in an HTML table. The binary is built on
# first use; OCR is serialised — it is CPU-bound and the build must happen once.

OCR_BUILD = threading.Lock()
OCR_GATE = threading.Lock()


def ocr_exe() -> Path:
    """Compile site/scripts/ocr_vision.swift into tmp/ocr/ocr_vision on first use."""
    with OCR_BUILD:
        if OCR_BIN.exists() and OCR_BIN.stat().st_mtime >= OCR_SRC.stat().st_mtime:
            return OCR_BIN
        OCR_DIR.mkdir(parents=True, exist_ok=True)
        r = subprocess.run(["swiftc", "-O", "-o", str(OCR_BIN), str(OCR_SRC)],
                           capture_output=True, text=True)
        if r.returncode != 0 or not OCR_BIN.exists():
            raise RuntimeError(f"swiftc failed: {(r.stderr or r.stdout).strip()[:200]}")
        return OCR_BIN


def ocr_images(paths, langs=OCR_LANGS) -> str:
    """Recognised text of each image, in the given order, blocks joined by a blank line."""
    if not paths:
        return ""
    exe = ocr_exe()
    r = subprocess.run([str(exe), "--langs", langs, *[str(p) for p in paths]],
                       capture_output=True, text=True, timeout=900)
    if r.returncode != 0:
        raise RuntimeError(f"ocr_vision failed: {(r.stderr or r.stdout).strip()[:200]}")
    out, keep = [], False
    for line in r.stdout.splitlines():
        if line.startswith("===== FILE "):
            keep = True
            out.append("")
        elif line.startswith("===== ERROR "):
            keep = False
        elif keep:
            out.append(line)
    return "\n".join(out).strip()


def ocr_pdf(path: Path, key: str, dpi=OCR_DPI, langs=OCR_LANGS) -> str:
    """Render the PDF's pages to bitmaps and OCR them."""
    import pypdfium2

    pdf = pypdfium2.PdfDocument(str(path))
    try:
        pages = min(len(pdf), OCR_MAX_PAGES)
        shots = []
        for i in range(pages):
            png = OCR_DIR / f"{key}_p{i + 1:03d}.png"
            pdf[i].render(scale=dpi / 72).to_pil().save(png)
            shots.append(png)
    finally:
        pdf.close()
    try:
        return ocr_images(shots, langs)
    finally:
        for p in shots:
            try:
                p.unlink()
            except OSError:
                pass


def ocr_source(kind: str, path: Path, key: str, dpi=OCR_DPI, langs=OCR_LANGS) -> str:
    OCR_DIR.mkdir(parents=True, exist_ok=True)
    with OCR_GATE:
        return ocr_pdf(path, key, dpi, langs) if kind == "pdf" else ocr_images([path], langs)


# ------------------------------------------------------------ non-HTML sources / fallbacks

def yandex_href(url: str, gate: "HostGate"):
    """Real download URL behind a Yandex Disk public link (the page itself is a JS shell)."""
    status, _ct, body, _final = gate.fetch(YANDEX_API + urllib.parse.quote(url, safe=""))
    if 200 <= status < 300:
        return json.loads(body.decode("utf-8", "replace")).get("href")
    return None


def wayback_timestamp(url: str, gate: "HostGate"):
    """Newest 200 snapshot of `url` in the Wayback Machine, or None. The CDX index is
    asked first (the availability API rate-limits hard), the availability API second."""
    q = urllib.parse.quote(url, safe="")
    try:
        status, _ct, body, _final = gate.fetch(WAYBACK_CDX.format(q), ARCHIVE_TIMEOUT)
        rows = json.loads(body or b"[]") if 200 <= status < 300 else []
        if len(rows) > 1:  # rows[0] is the header, the rest are ascending by timestamp
            return rows[-1][0]
    except Exception:
        pass
    try:
        status, _ct, body, _final = gate.fetch(WAYBACK_AVAIL.format(q), ARCHIVE_TIMEOUT)
        if 200 <= status < 300:
            snap = (json.loads(body or b"{}").get("archived_snapshots") or {}).get("closest") or {}
            if snap.get("available") and snap.get("timestamp"):
                return snap["timestamp"]
    except Exception:
        pass
    return None


def decode(body: bytes, ctype: str) -> str:
    encs = ["utf-8"]  # strict utf-8 first: mis-declared charsets are common, false utf-8 is not
    m = re.search(r"charset=[\"']?([\w-]+)", ctype or "", re.I)
    if m:
        encs.append(m.group(1))
    m = re.search(rb"charset=[\"']?([\w-]+)", body[:4096], re.I)
    if m:
        encs.append(m.group(1).decode("ascii", "ignore"))
    encs += ["cp1252", "latin-1"]
    for enc in encs:
        try:
            return body.decode(enc)
        except (LookupError, UnicodeDecodeError):
            continue
    return body.decode("latin-1", "replace")


def attr(tag: str, name: str) -> str:
    m = re.search(r'(?is)\b%s\s*=\s*(?:"([^"]*)"|\'([^\']*)\'|([^\s>]+))' % re.escape(name), tag)
    return (m.group(1) or m.group(2) or m.group(3) or "") if m else ""


def html_text(s: str) -> str:
    extra = []  # meta descriptions carry the caption/lede on news + social pages
    for tag in re.findall(r"(?is)<meta\b[^>]*>", s):
        key = (attr(tag, "name") or attr(tag, "property")).lower()
        if key in ("description", "og:description", "og:title", "twitter:description", "twitter:title"):
            extra.append(attr(tag, "content"))
    s = re.sub(r"(?is)<(script|style|noscript|svg|template)\b.*?</\1\s*>", " ", s)
    s = re.sub(r"(?s)<!--.*?-->", " ", s)
    s = re.sub(r"<[^>]+>", " ", s)
    return html.unescape(s + "\n" + "\n".join(extra))


def pdf_text(body: bytes) -> str:
    """pypdf first (fast); pdfplumber when that yields next to nothing (it keeps
    column order on some generators pypdf gives up on). Empty = a scan, OCR's job."""
    out = ""
    try:
        from pypdf import PdfReader

        out = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(body)).pages)
    except Exception:
        out = ""
    if len(out.strip()) >= MIN_TEXT:
        return out
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(body)) as pdf:
            alt = "\n".join((p.extract_text() or "") for p in pdf.pages)
        if len(alt.strip()) > len(out.strip()):
            out = alt
    except Exception:
        pass
    return out


def xlsx_text(body: bytes) -> str:
    """Every sheet, every non-empty row, one line each, cells joined with " | " so the
    column-split rules read a "Country | Name | score" table the way they read HTML."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    lines = []
    try:
        for ws in wb.worksheets:
            lines.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [("" if c is None else str(c)).strip() for c in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    lines.append(" | ".join(cells))
    finally:
        wb.close()
    return "\n".join(lines)


def extract(url: str, ctype: str, body: bytes):
    """-> (kind, text) with kind in html | pdf | xlsx | image | text | binary.
    `image` (and a `pdf` with no text layer) comes back empty here; ensure_fetched
    sends those through OCR."""
    ct = (ctype or "").lower()
    path = url.lower().split("?")[0].split("#")[0]
    if body[:5] == b"%PDF-" or "pdf" in ct or path.endswith(".pdf"):
        return "pdf", pdf_text(body)
    if path.endswith(SHEET_EXT) or any(x in ct for x in SHEET_CT):
        try:
            return "xlsx", xlsx_text(body)
        except Exception:
            return "binary", ""
    if any(x in ct for x in IMAGE_CT) or (path.endswith(IMAGE_EXT) and not ct.startswith("text/")):
        return "image", ""
    if body[:2] == b"PK" or any(x in ct for x in BINARY_CT):
        return "binary", ""
    s = decode(body, ctype)
    if "html" in ct or "xml" in ct or re.search(r"(?i)<(html|head|body|div|p|table|br)\b", s[:8192]):
        return "html", html_text(s)
    return "text", s


def cache_key(url: str) -> str:
    return hashlib.sha1(url.encode("utf-8")).hexdigest()


def cache_paths(url: str):
    """-> (meta json, extracted text, raw bytes of a binary kind)."""
    key = cache_key(url)
    return CACHE / f"{key}.json", CACHE / f"{key}.txt", CACHE / f"{key}.bin"


def atomic_write(path: Path, text: str):
    tmp = path.with_suffix(path.suffix + ".part")
    tmp.write_text(text, encoding="utf-8", errors="replace")
    tmp.replace(path)


def take(url: str, meta: dict, status: int, ctype: str, body: bytes, final: str) -> str:
    """Record one response in `meta` and return its text. Binary kinds keep their raw
    bytes next to the cache entry so OCR can be redone without asking the host again."""
    meta.update(http=status, contentType=ctype, finalUrl=final, bytes=len(body))
    meta["ok"] = 200 <= status < 300
    if not meta["ok"]:
        return ""
    text = ""
    try:
        meta["kind"], text = extract(url, ctype, body)
    except Exception as e:  # extraction bug/oddity is not a fetch failure
        meta["kind"], meta["extractError"] = "binary", f"{type(e).__name__}: {e}"[:200]
    bin_p = cache_paths(url)[2]
    if meta["kind"] in ("pdf", "xlsx", "image", "binary"):
        bin_p.write_bytes(body)
    elif bin_p.exists():
        bin_p.unlink()
    if meta["kind"] in ("pdf", "image") and len(text.strip()) < MIN_TEXT:
        try:
            text = ocr_source(meta["kind"], bin_p, cache_key(url))
            meta["ocr"] = True
        except Exception as e:
            meta["ocrError"] = f"{type(e).__name__}: {e}"[:200]
    return text


def ensure_fetched(url: str, gate: HostGate, retry_errors: bool, skip_hosts=frozenset(), force=frozenset()):
    """Fetch + extract once; -> (meta, fetched_now)."""
    meta_p, txt_p, _bin_p = cache_paths(url)
    if meta_p.exists() and txt_p.exists() and url not in force:
        meta = json.loads(meta_p.read_text(encoding="utf-8"))
        if not (retry_errors and not meta.get("ok")):
            return meta, False
    if urllib.parse.urlparse(url).netloc.lower() in skip_hosts:
        if meta_p.exists():
            return json.loads(meta_p.read_text(encoding="utf-8")), False
        meta = {"url": url, "ok": False, "error": "skipped (--skip-hosts)", "textChars": 0}
        atomic_write(txt_p, "")
        atomic_write(meta_p, json.dumps(meta, ensure_ascii=False))
        return meta, False
    meta = {"url": url, "fetchedAt": datetime.now(timezone.utc).isoformat(timespec="seconds")}
    text, target = "", url.split("#")[0]
    if urllib.parse.urlparse(url).netloc.lower() in YANDEX_HOSTS:
        try:  # the public page is a JS shell; the API hands out the file itself
            href = yandex_href(target, gate)
            if href:
                meta.update(via="yandex-disk", downloadUrl=href[:300])
                target = href
        except Exception as e:
            meta["resolveError"] = f"{type(e).__name__}: {e}"[:200]
    try:
        text = take(url, meta, *gate.fetch(target))
    except Exception as e:  # HTTPError, URLError, timeout, ssl, ...
        meta["ok"] = False
        meta["error"] = f"{type(e).__name__}: {e}"[:200]
    if not meta.get("ok"):  # dead host / 404 / broken TLS: the archive, never over a live answer
        try:
            ts = wayback_timestamp(target, gate)
            if ts:
                snap = WAYBACK_RAW.format(ts, target)
                live = meta.get("error")
                text = take(url, meta, *gate.fetch(snap, ARCHIVE_TIMEOUT))
                if meta["ok"]:
                    meta.pop("error", None)
                    meta.update(via="wayback", waybackTimestamp=ts, waybackUrl=snap, liveError=live)
                else:
                    meta["error"] = live
        except Exception as e:
            meta["waybackError"] = f"{type(e).__name__}: {e}"[:200]
    meta["textChars"] = len(text)
    atomic_write(txt_p, text)
    atomic_write(meta_p, json.dumps(meta, ensure_ascii=False))
    return meta, True


# ----------------------------------------------------------------------------- analysis

def resolve_in_edition(form: str, key, rosters):
    """The per-dataset graph builders also merge spellings that do not slugify alike
    ("Ivna de Gomes" into Ivna de Lima Ferreira Gomes). When the slug route misses, take
    the one roster member of that edition whose name tokens nest with the recorded
    form's; more than one candidate, or none, means give up rather than guess."""
    want = {t for t in tokens(form) if t not in PARTICLES}
    if len(want) < 2:
        return None
    hits = [pid for pid, specs in rosters.get(key, ())
            if want <= {t for sp in specs for t in sp.sig} or {t for sp in specs for t in sp.sig} <= want]
    return hits[0] if len(hits) == 1 else None


def raw_name_forms(rosters):
    """(olympiad, year, person id) -> the name form(s) the dataset itself recorded for
    that participation. people.json carries the display name and, unless the person is in
    `hideVariants`, the merged spellings — so the form a source of that edition actually
    prints can be missing from the pool entirely. It is read back here from the dataset's
    own raw records and resolved to the site person id through build_people's slugifier
    and slugAliases, the same chain build_people uses."""
    slug_aliases = json.loads((SITE / "scripts" / "aliases.json").read_text(encoding="utf-8")).get("slugAliases", {})
    members = {k: {pid for pid, _ in v} for k, v in rosters.items()}
    extra, unresolved = defaultdict(list), []
    for ds in DATASETS:
        raw = ROOT / ds / "data" / "raw" / f"{ds}.json"
        if not raw.exists():
            continue
        for rec in json.loads(raw.read_text(encoding="utf-8")):
            form, key = rec.get("rawName"), (ds, rec.get("year"))
            if not form:
                continue
            slug = slugify(form)
            pid = slug_aliases.get(slug, slug)
            if pid not in members.get(key, ()):
                pid = resolve_in_edition(form, key, rosters)
            if pid is None:
                unresolved.append((ds, key[1], form))
            elif form not in extra[(key[0], key[1], pid)]:
                extra[(key[0], key[1], pid)].append(form)
    return extra, unresolved


def load_rosters():
    people = json.loads(PEOPLE.read_text(encoding="utf-8"))["people"]
    rosters = defaultdict(list)  # (ol, year) -> [(pid, [NameSpec])]
    parts = []  # (pid, display name, ol, year)
    for p in people:
        specs, seen = [], set()
        for form in [p["name"]] + list(p.get("nameVariants", [])):
            sp = NameSpec(form)
            if sp.key not in seen:
                seen.add(sp.key)
                specs.append(sp)
        for x in p["participations"]:
            rosters[(x["olympiad"], x["year"])].append((p["id"], specs))
            parts.append((p["id"], p["name"], x["olympiad"], x["year"]))
    # the recorded form is evidence about one edition, not about the person: it is added
    # to that (olympiad, year) roster entry only, never to the person's pool everywhere
    extra, unresolved = raw_name_forms(rosters)
    widened = 0
    for key, members in list(rosters.items()):
        rebuilt = []
        for pid, specs in members:
            seen, grown = {sp.key for sp in specs}, list(specs)
            for form in extra.get((key[0], key[1], pid), ()):
                sp = NameSpec(form)
                if sp.key not in seen:
                    seen.add(sp.key)
                    grown.append(sp)
            if len(grown) > len(specs):
                widened += 1
                specs = grown
            rebuilt.append((pid, specs))
        rosters[key] = rebuilt
    print(f"rosters: {widened} participation(s) widened by a recorded form the pool lacks"
          + (f"; {len(unresolved)} raw record(s) unmatched" if unresolved else ""), flush=True)
    return rosters, parts


def classify(url: str, meta: dict, doc):
    """A spreadsheet read cell by cell and a scan read by OCR are `text` like any HTML
    page — what decides is whether enough folded text came out of the source at all."""
    if not meta.get("ok"):
        return "error"
    host = urllib.parse.urlparse(url).netloc.lower()
    if any(host == h or host.endswith("." + h) for h in NO_TEXT_HOSTS):
        return "status-only"
    if meta.get("kind") == "binary" or doc is None or len(doc.text) < MIN_TEXT:
        return "status-only"
    return "text"


def cached_status(url: str, meta: dict, raw: str) -> str:
    """classify() straight off a cache entry (a JS shell can be long and still fold to
    nothing, so the fold has to happen — no length shortcut)."""
    return classify(url, meta, Doc(raw) if raw.strip() else None)


def analyse(datasets, corr, rosters, metas):
    docs = {}

    def doc_for(u):
        if u not in docs:
            txt_p = cache_paths(u)[1]
            raw = txt_p.read_text(encoding="utf-8", errors="replace") if txt_p.exists() else ""
            docs[u] = Doc(raw) if raw.strip() else None
        return docs[u]

    by_ol = {}
    expansions = {}  # pid -> (longer form seen on a page, url, ds, year)
    variants = []  # (ds, year, pid, url, raw span) for every near-miss match, for review
    shorts = []    # the same for every shortened-form match
    bad_attested = []  # "names" ids that are not on that edition's roster (typos)
    for ds in datasets:
        # every other person of the dataset, for the "is this somebody else?" guard
        ds_own = defaultdict(set)
        for (o, _y), r in rosters.items():
            if o == ds:
                for pid, specs in r:
                    ds_own[pid].update(t for sp in specs for t in sp.sig)
        pool_for = {pid: tuple(v for k, v in ds_own.items() if k != pid) for pid in ds_own}
        by_year = {}
        for yrec in corr[ds]["years"]:
            year = yrec["year"]
            roster = rosters.get((ds, year), [])
            all_tokens = {t for _, specs in roster for sp in specs for t in sp.core if len(t) >= 3}
            own_sets = {pid: {t for sp in specs for t in sp.sig} for pid, specs in roster}
            by_url = {}
            for s in yrec["sources"]:
                url = s["url"]
                doc = doc_for(url)
                status = classify(url, metas[url], doc)
                found = {}
                if status == "text":
                    for pid, specs in roster:
                        own = own_sets[pid]
                        others = all_tokens - own
                        mates = tuple(v for k, v in own_sets.items() if k != pid)
                        best, span = None, None
                        for sp in specs:
                            lv, longer, vspan = match_level(doc, sp, others, mates, pool_for[pid])
                            if lv and RANK[lv] > RANK.get(best, 0):
                                best, span = lv, vspan
                            if longer:
                                extra = [t for t in tokens(longer) if len(t) >= 3 and t not in own and t not in PARTICLES]
                                if extra and (pid not in expansions or len(longer) > len(expansions[pid][0])):
                                    expansions[pid] = (longer, url, ds, year)
                        if best:
                            found[pid] = best
                            if best in ("variant", "short"):
                                (variants if best == "variant" else shorts).append((ds, year, pid, url, span))
                # curated attestation: the source names these people even though its
                # text could not be read here; never overrides a match found above
                members = {pid for pid, _ in roster}
                for pid in s.get("names", ()):
                    if pid in members:
                        found.setdefault(pid, "attested")
                    else:
                        bad_attested.append((ds, year, pid, url))
                by_url[url] = {"status": status, "found": dict(sorted(found.items()))}
            by_year[str(year)] = by_url
        by_ol[ds] = by_year
    for ds, year, pid, url in bad_attested:
        print(f"WARNING: {ds} {year} attests `{pid}` on {url}, but that person is not on the edition's roster", flush=True)
    return by_ol, expansions, variants, shorts


def cell(s: str) -> str:
    """One markdown table cell: page text may itself contain '|'."""
    return re.sub(r"\s+", " ", str(s)).replace("|", "\\|").strip()


def write_report(by_ol, rosters, parts, generated_at, metas, expansions, variants, shorts):
    L = ["# Source name-presence report", "", f"Generated {generated_at} by `site/scripts/check_source_names.py`.", ""]
    L += [
        "A (url, person) pair is **found** when the page text names the roster member",
        "(levels: full > firstlast > variant > surname > short, plus **attested** — a source",
        "whose text cannot be read, curated with `\"names\"`), **not found** when the page has",
        "text but no match, **unverifiable** when the page yielded no usable text and attests",
        "nothing (error / status-only).",
        "",
        "## Per dataset",
        "",
        f"| dataset | url refs | text | status-only | error | pairs found | not found | unverifiable | {' / '.join(LEVELS)} |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---|",
    ]
    tot = defaultdict(int)
    for ds in DATASETS:
        if ds not in by_ol:
            continue
        c = defaultdict(int)
        for year, by_url in by_ol[ds].items():
            roster_n = len(rosters.get((ds, int(year)), []))
            for url, rec in by_url.items():
                c[rec["status"]] += 1
                c["refs"] += 1
                c["found"] += len(rec["found"])
                for lv in rec["found"].values():
                    c[lv] += 1
                # an unreadable source is unverifiable for everyone it does not attest
                key = "notfound" if rec["status"] == "text" else "unverifiable"
                c[key] += roster_n - len(rec["found"])
        for k, v in c.items():
            tot[k] += v
        L.append(
            f"| {ds} | {c['refs']} | {c['text']} | {c['status-only']} | {c['error']} | {c['found']} | "
            f"{c['notfound']} | {c['unverifiable']} | {' / '.join(str(c[lv]) for lv in LEVELS)} |"
        )
    L.append(
        f"| **total** | {tot['refs']} | {tot['text']} | {tot['status-only']} | {tot['error']} | {tot['found']} | "
        f"{tot['notfound']} | {tot['unverifiable']} | {' / '.join(str(tot[lv]) for lv in LEVELS)} |"
    )

    # gap list: participations with no source that names the person
    gaps = []
    for pid, name, ol, year in parts:
        if ol not in by_ol:
            continue
        by_url = by_ol[ol].get(str(year))
        if by_url is None:
            gaps.append((ol, year, name, pid, "no sources for this edition"))
            continue
        if not any(pid in rec["found"] for rec in by_url.values()):
            n_unv = sum(1 for rec in by_url.values() if rec["status"] != "text")
            why = f"{len(by_url)} source(s), none names the person" + (f"; {n_unv} unverifiable" if n_unv else "")
            gaps.append((ol, year, name, pid, why))
    gaps.sort(key=lambda g: (g[0], g[1], g[2]))
    L += ["", f"## Participations with no name-bearing source ({len(gaps)} of {sum(1 for p in parts if p[2] in by_ol)})", ""]
    cur = None
    for ol, year, name, pid, why in gaps:
        if ol != cur:
            cur = ol
            L.append(f"### {ol} ({sum(1 for g in gaps if g[0] == ol)})")
        L.append(f"- {year} — {name} (`{pid}`) — {why}")

    # per-row triage of the gap list, if it has been recorded next to the report
    notes = json.loads(NOTES.read_text(encoding="utf-8")) if NOTES.exists() else {}
    if notes:
        L += ["", "## Gap triage", "", "| dataset | year | person | class | note |", "|---|---:|---|---|---|"]
        for ol, year, name, pid, _ in gaps:
            cls, note = notes.get(f"{ol} {year} {pid}", ["?", "not triaged"])
            L.append(f"| {ol} | {year} | {cell(name)} | {cls} | {cell(note)} |")

    # every near-miss match, with the page's own spelling, so a human can check it
    L += ["", f"## Variant (near-miss) matches ({len(variants)})", "",
          "One token one edit away from the pool form; everything else matched exactly.", "",
          "| dataset | year | person | page spelling | source |", "|---|---:|---|---|---|"]
    for ds, year, pid, url, span in sorted(variants, key=lambda v: (DATASETS.index(v[0]), v[1], v[2])):
        L.append(f"| {ds} | {year} | `{pid}` | {cell(span)} | {url} |")

    # every pair whose best level is a shortened form, with the page's own wording
    L += ["", f"## Short-form matches ({len(shorts)})", "",
          "The page prints only part of the name; no teammate of that edition fits the span.", "",
          "| dataset | year | person | page wording | host |", "|---|---:|---|---|---|"]
    for ds, year, pid, url, span in sorted(shorts, key=lambda v: (DATASETS.index(v[0]), v[1], v[2])):
        host = re.sub(r"^https?://(www\.)?", "", url).split("/")[0]
        L.append(f"| {ds} | {year} | `{pid}` | {cell(span)} | {host} |")

    # sources that needed more than a plain GET to read: spreadsheets, OCR'd scans and
    # photos, a resolved file-host link, an archived copy of a dead URL
    opened = [m for m in metas.values() if m.get("ok") and (m.get("ocr") or m.get("via") or m.get("kind") == "xlsx")]
    L += ["", f"## Sources opened beyond plain HTML ({len(opened)})", "",
          "| how | kind | chars | source |", "|---|---|---:|---|"]
    for m in sorted(opened, key=lambda m: (m.get("via") or ("ocr" if m.get("ocr") else "xlsx"), m["url"])):
        how = m.get("via") or ("ocr" if m.get("ocr") else "spreadsheet")
        if m.get("via") == "wayback":
            how += f" {m.get('waybackTimestamp', '')}"
        if m.get("ocr") and m.get("via"):
            how += " + ocr"
        L.append(f"| {how} | {m.get('kind', '-')} | {m.get('textChars', 0)} | {m['url']} |")

    # fetch failures by host: a host that fails wholesale is blocking us or gone,
    # and every pair behind it is "unverifiable" until --retry-errors succeeds
    fails = defaultdict(list)
    for m in metas.values():
        if not m.get("ok"):
            fails[urllib.parse.urlparse(m["url"]).netloc.lower()].append(m.get("error") or f"HTTP {m.get('http')}")
    L += ["", f"## Fetch failures by host ({sum(len(v) for v in fails.values())} unique URLs)", "",
          "| host | failed URLs | typical error |", "|---|---:|---|"]
    for host, errs in sorted(fails.items(), key=lambda kv: -len(kv[1])):
        typical = max(set(errs), key=errs.count)[:70]
        L.append(f"| {host} | {len(errs)} | {typical} |")

    # name-expansion candidates: a page shows first + middle names the pool lacks + last.
    # Ids curated into aliases.json `expansionsRejected` were looked at and turned down
    # (source typo, spelling-only difference): drop them so they stop being re-proposed.
    rejected = json.loads((SITE / "scripts" / "aliases.json").read_text(encoding="utf-8")).get("expansionsRejected", {})
    rows = sorted(((pid, v) for pid, v in expansions.items() if pid not in rejected),
                  key=lambda kv: (DATASETS.index(kv[1][2]), kv[1][3], kv[0]))
    skipped = sum(1 for pid in expansions if pid in rejected)
    L += ["", f"## Name-expansion candidates ({len(rows)} people: a page shows a longer form than the pool has"
          + (f"; {skipped} rejected earlier, see aliases.json)" if skipped else ")"), ""]
    for pid, (longer, url, ds, year) in rows:  # listed in full: the list is short and every row is a decision
        L.append(f"- {ds} {year} — `{pid}` → \"{longer}\" — {url}")

    # appendix: pages with text that name nobody on the roster (JS shells, soft-404s, counts-only)
    zero = []
    for ds in DATASETS:
        for year, by_url in by_ol.get(ds, {}).items():
            for url, rec in by_url.items():
                if rec["status"] == "text" and not rec["found"] and rosters.get((ds, int(year))):
                    zero.append((ds, year, url))
    L += ["", f"## Appendix: pages with text but no roster name for that edition ({len(zero)})", ""]
    L += [f"- {ds} {year} {url}" for ds, year, url in zero[:400]]
    if len(zero) > 400:
        L.append(f"- … {len(zero) - 400} more")
    REPORT.write_text("\n".join(L) + "\n", encoding="utf-8")
    return gaps, tot


# ----------------------------------------------------------------------------- main

def select_refetch(urls, spec: str) -> frozenset:
    """URLs whose cached entry matches one of the comma-separated statuses/kinds in
    `spec` — those get fetched again this run, everything else stays cache-only.
    Social hosts are excluded: an anonymous fetch of those never returns the post."""
    want = {s.strip().lower() for s in spec.split(",") if s.strip()}
    if not want:
        return frozenset()
    picked = set()
    for u in urls:
        host = urllib.parse.urlparse(u).netloc.lower()
        if any(host == h or host.endswith("." + h) for h in NO_TEXT_HOSTS):
            continue
        meta_p, txt_p, _ = cache_paths(u)
        if not meta_p.exists():
            continue
        meta = json.loads(meta_p.read_text(encoding="utf-8"))
        raw = txt_p.read_text(encoding="utf-8", errors="replace") if txt_p.exists() else ""
        if cached_status(u, meta, raw) in want or (meta.get("kind") or "") in want:
            picked.add(u)
    return frozenset(picked)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--only", help="comma-separated dataset ids (others are kept from the existing output)")
    ap.add_argument("--retry-errors", action="store_true", help="refetch URLs whose cached fetch failed")
    ap.add_argument("--refetch", default="",
                    help="comma-separated cache statuses (text, status-only, error) and/or kinds "
                         "(html, pdf, xlsx, image, binary, text) to invalidate and fetch again this "
                         "run; everything else stays cache-only. Social hosts are never included.")
    ap.add_argument("--skip-hosts", default="", help="comma-separated hosts not to (re)fetch this run, e.g. one that is rate-limiting us")
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--delay", type=float, default=1.0, help="seconds between requests to the same host")
    args = ap.parse_args()
    skip_hosts = {h.strip().lower() for h in args.skip_hosts.split(",") if h.strip()}
    datasets = args.only.split(",") if args.only else list(DATASETS)
    unknown = [d for d in datasets if d not in DATASETS]
    if unknown:
        sys.exit(f"unknown dataset(s): {unknown}")
    CACHE.mkdir(parents=True, exist_ok=True)

    corr = {ds: json.loads((ROOT / ds / "data" / "corroboration.json").read_text(encoding="utf-8")) for ds in datasets}
    urls = sorted({s["url"] for c in corr.values() for y in c["years"] for s in y["sources"]})
    by_host = defaultdict(list)
    for u in urls:
        by_host[urllib.parse.urlparse(u).netloc.lower()].append(u)
    groups = sorted(by_host.values(), key=len, reverse=True)
    order = [g[i] for i in range(len(groups[0])) for g in groups if i < len(g)]  # interleave hosts
    print(f"{len(order)} unique URLs on {len(by_host)} hosts ({sum(len(y['sources']) for c in corr.values() for y in c['years'])} refs)", flush=True)

    force = select_refetch(order, args.refetch)
    if args.refetch:
        print(f"--refetch {args.refetch}: {len(force)} cached URLs invalidated", flush=True)

    gate = HostGate(args.delay)
    metas, n_new, t0 = {}, 0, time.monotonic()
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs = {ex.submit(ensure_fetched, u, gate, args.retry_errors, skip_hosts, force): u for u in order}
        for k, fut in enumerate(as_completed(futs), 1):
            u = futs[fut]
            meta, fetched = fut.result()
            metas[u] = meta
            n_new += fetched
            if fetched or k % 100 == 0:
                tag = "GET" if fetched else "hit"
                st = "ok " if meta.get("ok") else "ERR"
                print(f"[{k}/{len(order)}] {tag} {st} {meta.get('kind', '-'):6} {meta.get('textChars', 0):>7}c {u[:100]}", flush=True)
    print(f"fetched {n_new}, cached {len(order) - n_new}, {time.monotonic() - t0:.0f}s", flush=True)

    rosters, parts = load_rosters()
    by_ol, expansions, variants, shorts = analyse(datasets, corr, rosters, metas)
    if args.only and OUT.exists():
        prev = json.loads(OUT.read_text(encoding="utf-8")).get("byOlympiad", {})
        prev.update(by_ol)
        by_ol = {ds: prev[ds] for ds in DATASETS if ds in prev}
    generated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    OUT.write_text(json.dumps({"generatedAt": generated_at, "byOlympiad": by_ol}, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    gaps, tot = write_report(by_ol, rosters, parts, generated_at, metas, expansions, variants, shorts)
    print(
        f"source_names.json: {tot['refs']} url refs — text {tot['text']}, status-only {tot['status-only']}, error {tot['error']}; "
        f"pairs found {tot['found']} (" + ", ".join(f"{lv} {tot[lv]}" for lv in LEVELS) + "), "
        f"not found {tot['notfound']}, unverifiable {tot['unverifiable']}"
    )
    print(f"gap list: {len(gaps)} participations without a name-bearing source -> {REPORT}")


if __name__ == "__main__":
    main()
